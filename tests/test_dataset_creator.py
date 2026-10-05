import io
import random
import tempfile
from datetime import datetime, time, timedelta
from pathlib import Path
import unittest
from contextlib import redirect_stdout
from unittest.mock import patch

from openpyxl import load_workbook

from dataset_creator import (
    available_products_for_store,
    create_bank_card,
    create_transaction_receipt,
    store_price_multiplier,
    MAX_CARD_USES,
    load_config,
    main,
)


class DatasetCreatorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = load_config()

    def test_every_store_has_products_with_configured_brands(self):
        for store in self.config["stores"]:
            with self.subTest(store=store["name"]):
                self.assertTrue(available_products_for_store(store, self.config))

    def test_every_store_has_a_valid_synthetic_price_tier(self):
        tiers = self.config["store_price_model"]["store_tiers"]
        multipliers = self.config["store_price_model"]["tier_multipliers"]
        self.assertEqual(
            set(tiers),
            {store["name"] for store in self.config["stores"]},
        )
        self.assertEqual(set(multipliers), {"discount", "standard", "premium"})
        for store in self.config["stores"]:
            with self.subTest(store=store["name"]):
                self.assertGreater(store_price_multiplier(store, self.config), 0)

    def test_transaction_category_and_brand_fit_selected_store(self):
        for _ in range(50):
            receipt = create_transaction_receipt(self.config)
            store = next(
                store
                for store in self.config["stores"]
                if store["name"] == receipt["store_name"]
            )

            available_categories = available_products_for_store(store, self.config)
            for item in receipt["products"]:
                self.assertIn(item["category"], available_categories)
                matching_brand = next(
                    brand
                    for brand in self.config["brands"][item["category"]]
                    if (brand["name"] if isinstance(brand, dict) else str(brand))
                    == item["brand"]
                )
                if isinstance(matching_brand, dict):
                    min_price = int(
                        matching_brand.get("min_price")
                        or matching_brand.get("price_min")
                        or matching_brand.get("min")
                        or 1000
                    )
                    max_price = int(
                        matching_brand.get("max_price")
                        or matching_brand.get("price_max")
                        or matching_brand.get("max")
                        or max(min_price, 10000)
                    )
                else:
                    min_price, max_price = 1000, 10000
                multiplier = store_price_multiplier(store, self.config)
                self.assertGreaterEqual(
                    item["item_price"], max(1, round(min_price * multiplier))
                )
                self.assertLessEqual(
                    item["item_price"],
                    max(
                        max(1, round(min_price * multiplier)),
                        round(max_price * multiplier),
                    ),
                )
            self.assertGreaterEqual(
                sum(item["quantity"] for item in receipt["products"]), 2
            )
            self.assertLessEqual(
                sum(item["quantity"] for item in receipt["products"]), 50
            )
            self.assertEqual(
                receipt["total_amount"],
                sum(
                    item["item_price"] * item["quantity"]
                    for item in receipt["products"]
                ),
            )
            purchase_time = datetime.fromisoformat(receipt["timestamp"]).replace(tzinfo=None)
            latitude, longitude = receipt["coords"].split(",")
            branch = next(
                branch
                for branch in store["branches"]
                if f"{branch['lat']:.8f}" == latitude
                and f"{branch['lon']:.8f}" == longitude
            )
            self.assertEqual(purchase_time.year, 2026)
            opening = datetime.combine(purchase_time.date(), time.fromisoformat(branch["open"]))
            closing = datetime.combine(purchase_time.date(), time.fromisoformat(branch["close"]))
            if closing <= opening:
                closing += timedelta(days=1)
            self.assertGreaterEqual(purchase_time, opening)
            self.assertLess(purchase_time, closing)
            self.assertAlmostEqual(float(latitude), branch["lat"], places=8)
            self.assertAlmostEqual(float(longitude), branch["lon"], places=8)

    def test_main_writes_receipt_items_as_rows_in_xlsx(self):
        columns = [
            "Название магазина",
            "Дата и время",
            "Координаты",
            "Категория",
            "Бренд",
            "Цена товара, руб.",
            "Количество, шт.",
            "Номер карты",
            "Номер чека",
            "Итого по чеку, руб.",
        ]
        receipt = {
            "store_name": "Тестовый магазин",
            "timestamp": "2026-01-01T10:00+03:00",
            "coords": "59.90000000,30.30000000",
            "products": [
                {"category": "ноутбук", "brand": "Lenovo", "item_price": 50000, "quantity": 1},
                {"category": "смартфон", "brand": "Xiaomi", "item_price": 20000, "quantity": 1},
            ],
            "card_number": "1234 5678 1234 5678",
            "receipt_number": "№ 1",
            "total_amount": 70000,
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "transactions.xlsx"
            output = io.StringIO()
            with (
                patch("dataset_creator.OUTPUT_PATH", output_path),
                patch("dataset_creator.NUM_RECEIPTS", 1),
                patch("dataset_creator.create_transaction_receipt", return_value=receipt),
                redirect_stdout(output),
            ):
                main()

            workbook = load_workbook(output_path, data_only=True)
            worksheet = workbook["Чеки"]

        self.assertEqual([cell.value for cell in worksheet[1]], columns)
        self.assertEqual(worksheet.max_row, 3)
        self.assertEqual(
            [worksheet.cell(row, 4).value for row in (2, 3)],
            ["ноутбук", "смартфон"],
        )
        self.assertEqual(worksheet["A2"].value, "Тестовый магазин")
        self.assertIsNone(worksheet["A3"].value)
        self.assertEqual(worksheet["J2"].value, 70000)
        for column in ("B", "C", "H", "I", "J"):
            self.assertIsNone(worksheet[f"{column}3"].value)
        self.assertFalse(worksheet.merged_cells.ranges)
        self.assertIn("Филиалов:", output.getvalue())

    def test_generated_card_uses_configured_bank_and_payment_system(self):
        card = create_bank_card(self.config)
        bank = next(
            bank for bank in self.config["banks"]
            if bank["name"] == card["bank"]
        )

        self.assertEqual(len(card["number"]), 16)
        self.assertTrue(card["number"].isdigit())
        self.assertIn(card["payment_system"], bank["payment_systems"])
        self.assertTrue(any(
            card["number"].startswith(prefix)
            for prefix in bank["payment_systems"][card["payment_system"]]["bins"]
        ))

    def test_card_usage_is_shared_across_receipts_and_stops_at_five(self):
        card = create_bank_card(self.config)

        for expected_uses in range(1, MAX_CARD_USES + 1):
            create_transaction_receipt(self.config, card)
            self.assertEqual(card["used"], expected_uses)

        with self.assertRaisesRegex(ValueError, "уже использована 5 раз"):
            create_transaction_receipt(self.config, card)
        self.assertEqual(card["used"], MAX_CARD_USES)

    def test_receipt_numbers_are_unique_within_a_branch(self):
        card = create_bank_card(self.config)
        used_receipt_numbers = {}
        receipt_numbers = iter((123, 123, 124))
        original_randint = random.randint

        def controlled_randint(start, end):
            if (start, end) == (1, 999999):
                return next(receipt_numbers)
            return original_randint(start, end)

        with (
            patch("dataset_creator.random.choice", side_effect=lambda values: values[0]),
            patch("dataset_creator.random.randint", side_effect=controlled_randint),
        ):
            first = create_transaction_receipt(
                self.config, card, used_receipt_numbers
            )
            second = create_transaction_receipt(
                self.config, card, used_receipt_numbers
            )

        self.assertEqual(first["store_name"], second["store_name"])
        self.assertEqual(first["coords"], second["coords"])
        self.assertEqual(first["receipt_number"], "№ 123")
        self.assertEqual(second["receipt_number"], "№ 124")
        self.assertEqual(len(used_receipt_numbers), 1)

    def test_receipt_can_contain_multiple_distinct_products(self):
        store = next(
            store for store in self.config["stores"]
            if len(available_products_for_store(store, self.config)) > 1
        )
        config = {**self.config, "stores": [store]}
        card = create_bank_card(self.config)
        original_randint = random.randint

        def controlled_randint(start, end):
            if (start, end) == (2, 50):
                return 2
            if start == 1 and 2 <= end <= 5:
                return 2
            return original_randint(start, end)

        with patch("dataset_creator.random.randint", side_effect=controlled_randint):
            receipt = create_transaction_receipt(config, card)

        self.assertEqual(len(receipt["products"]), 2)
        product_pairs = {
            (item["category"], item["brand"]) for item in receipt["products"]
        }
        self.assertEqual(len(product_pairs), 2)
        self.assertEqual(
            len({item["category"] for item in receipt["products"]}), 2
        )
        self.assertEqual(sum(item["quantity"] for item in receipt["products"]), 2)

    def test_main_never_uses_a_card_more_than_five_times(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "transactions.xlsx"
            output = io.StringIO()
            with (
                patch("dataset_creator.OUTPUT_PATH", output_path),
                patch("dataset_creator.NUM_RECEIPTS", 10),
                redirect_stdout(output),
            ):
                main()

            workbook = load_workbook(output_path, data_only=True)
            worksheet = workbook["Чеки"]

        receipt_rows = [
            row for row in range(2, worksheet.max_row + 1)
            if worksheet.cell(row, 8).value is not None
        ]
        self.assertEqual(len(receipt_rows), 10)
        uses_by_card = {}
        for row in receipt_rows:
            card_number = worksheet.cell(row, 8).value
            uses_by_card[card_number] = uses_by_card.get(card_number, 0) + 1
        self.assertTrue(all(uses <= MAX_CARD_USES for uses in uses_by_card.values()))
        self.assertEqual(sum(uses_by_card.values()), 10)


if __name__ == "__main__":
    unittest.main()
